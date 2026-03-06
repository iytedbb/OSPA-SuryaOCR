
import os
import json
import uuid
import zipfile
import mimetypes
from io import BytesIO
from pathlib import Path
from datetime import datetime
from flask import render_template, request, jsonify, send_from_directory, send_file, Response, current_app
from sqlalchemy import text
from werkzeug.utils import secure_filename

from app.extensions import db
from app.modules.common.utils import secure_filename_tr, generate_detailed_filename, extract_year_smart, resolve_file_path
from . import admin_bp

# Helper to close db session is handled by Flask's teardown_request automatically when using SQLAlchemy
# However, explicit close isn't harmful but redundant. We will use db.session directly.

@admin_bp.route('/')
def index():
    """Ana sayfa"""
    return render_template('admin/index.html')

@admin_bp.route('/favicon.ico')
def favicon():
    # Admin static folder or global static? Let's assume global static/admin/gemini...
    return send_from_directory(
        os.path.join(current_app.root_path, 'static', 'admin'),
        'gemini_ospa_ocr_db_logo.png', mimetype='image/vnd.microsoft.icon'
    )

@admin_bp.route('/api/health', methods=['GET'])
def health_check():
    """Sistem sağlık kontrolü"""
    try:
        db.session.execute(text("SELECT 1"))
        db_status = "connected"
    except Exception as e:
        db_status = f"error: {str(e)}"
    
    return jsonify({
        "status": "ok",
        "database": db_status,
        "timestamp": datetime.now().isoformat()
    })

# ==========================================
# STATISTICS
# ==========================================

@admin_bp.route('/api/stats', methods=['GET'])
def get_statistics():
    """Dashboard istatistiklerini döndür"""
    print("DEBUG: /api/stats request received")
    try:
        stats = {}
        
        # 1. Toplam yazar sayısı
        result = db.session.execute(text(
            "SELECT COUNT(DISTINCT author) FROM documents WHERE author IS NOT NULL AND author != ''"
        ))
        stats['total_authors'] = result.scalar() or 0
        
        # 2. Toplam eser sayısı
        result = db.session.execute(text("SELECT COUNT(*) FROM documents"))
        stats['total_documents'] = result.scalar() or 0
        
        # 3. OCR tamamlanan
        try:
            result = db.session.execute(text("SELECT COUNT(*) FROM ocr_results"))
            stats['ocr_completed'] = result.scalar() or 0
        except Exception:
            stats['ocr_completed'] = 0
            
        # 4. Toplam sayfa
        try:
            result = db.session.execute(text("SELECT COALESCE(SUM(total_pages), 0) FROM ocr_results"))
            stats['total_pages'] = int(result.scalar() or 0)
        except Exception:
            stats['total_pages'] = 0
            
        # 5. Toplam dosya boyutu (MB)
        try:
            result = db.session.execute(text("SELECT COALESCE(SUM(file_size), 0) FROM document_files"))
            total_bytes = result.scalar() or 0
            stats['total_size_mb'] = round(float(total_bytes) / 1024 / 1024, 2)
        except Exception:
            stats['total_size_mb'] = 0
            
        # 6. Son 7 günde eklenen
        from datetime import timedelta
        # Ensure we are using a naive or aware datetime consistent with DB
        seven_days_ago = datetime.now() - timedelta(days=7)
        try:
            result = db.session.execute(text("""
                SELECT COUNT(*) FROM documents 
                WHERE created_at >= :date
            """), {"date": seven_days_ago})
            stats['recent_documents'] = result.scalar() or 0
        except Exception as e:
            print(f"DEBUG: recent_documents query failed: {e}")
            stats['recent_documents'] = 0
            
        # 7. En çok eser sahibi yazarlar (top 5)
        result = db.session.execute(text("""
            SELECT author, COUNT(*) as count
            FROM documents
            WHERE author IS NOT NULL AND author != ''
            GROUP BY author
            ORDER BY count DESC
            LIMIT 5
        """))
        stats['top_authors'] = [
            {"name": str(row[0]), "count": int(row[1])} 
            for row in result.fetchall()
        ]
        
        # 8. Metadata türlerine göre dağılım
        result = db.session.execute(text("""
            SELECT COALESCE(metadata_type, 'unknown') as type, COUNT(*) as count
            FROM documents
            GROUP BY metadata_type
            ORDER BY count DESC
        """))
        stats['type_distribution'] = [
            {"type": str(row[0]), "count": int(row[1])} 
            for row in result.fetchall()
        ]
        
        # 9. Son eklenen 5 eser
        result = db.session.execute(text("""
            SELECT id, title, author, created_at
            FROM documents
            ORDER BY created_at DESC
            LIMIT 5
        """))
        
        recent_additions = []
        for row in result.fetchall():
            created_at = row[3]
            # Handle possible string format from SQLite
            if created_at and isinstance(created_at, str):
                try:
                    created_at = datetime.fromisoformat(created_at.replace('Z', '+00:00'))
                except:
                    pass
            
            recent_additions.append({
                "id": str(row[0]),
                "title": str(row[1]) if row[1] else "Başlıksız",
                "author": str(row[2]) if row[2] else "Bilinmeyen",
                "created_at": created_at.isoformat() if created_at and hasattr(created_at, 'isoformat') else str(created_at) if created_at else None
            })
        stats['recent_additions'] = recent_additions
        
        print("DEBUG: /api/stats success")
        return jsonify(stats)
        
    except Exception as e:
        import traceback
        print(f"Stats Error: {e}")
        print(traceback.format_exc())
        return jsonify({
            "error": str(e), 
            "traceback": traceback.format_exc(),
            "status": "error"
        }), 500

# ==========================================
# AUTHORS
# ==========================================

@admin_bp.route('/api/authors', methods=['GET'])
def get_authors():
    """Yazarları listele"""
    try:
        sort = request.args.get('sort', 'a_z')
        search = request.args.get('search', '').strip()
        
        # Sıralama
        order_clause = "author ASC"
        if sort == 'z_a':
            order_clause = "author DESC"
        elif sort == 'doc_count':
            order_clause = "doc_count DESC"
        elif sort == 'recent':
            order_clause = "last_updated DESC"
        
        # Arama filtresi
        where_clause = "WHERE author IS NOT NULL AND author != ''"
        params = {}
        
        if search:
            # SQLite'da LIKE case-insensitive'dir (ASCII için), Postgres'te ILIKE gerekir.
            # db.engine.name ile kontrol edebiliriz ama LIKE çoğu durumda yeterlidir.
            like_op = "ILIKE" if db.engine.name == 'postgresql' else "LIKE"
            where_clause += f" AND author {like_op} :search"
            params['search'] = f"%{search.lower()}%"
        
        query = f"""
            SELECT 
                author,
                COUNT(*) as doc_count,
                MAX(created_at) as last_updated
            FROM documents
            {where_clause}
            GROUP BY author
            ORDER BY {order_clause}
        """
        
        result = db.session.execute(text(query), params)
        authors = [
            {
                "name": row[0],
                "doc_count": row[1],
                "last_updated": row[2].isoformat() if row[2] else None
            }
            for row in result.fetchall()
        ]
        
        return jsonify(authors)
        
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/authors/<author_name>', methods=['PUT'])
def update_author(author_name):
    """Yazar adını güncelle (tüm belgelerinde)"""
    try:
        data = request.get_json()
        new_name = data.get('new_name', '').strip()
        
        if not new_name:
            return jsonify({"error": "Yeni isim boş olamaz"}), 400
        
        result = db.session.execute(text("""
            UPDATE documents
            SET author = :new_name, updated_at = :now
            WHERE author = :old_name
        """), {"old_name": author_name, "new_name": new_name, "now": datetime.now()})
        
        db.session.commit()
        
        return jsonify({
            "success": True,
            "updated_count": result.rowcount,
            "message": f"'{author_name}' -> '{new_name}' olarak güncellendi"
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500

# ==========================================
# DOCUMENTS (RECORDS)
# ==========================================

@admin_bp.route('/api/records', methods=['GET'])
def get_records():
    """Kayıtları listele"""
    try:
        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 25))
        sort = request.args.get('sort', 'a_z')
        author = request.args.get('author', '').strip()
        search = request.args.get('search', '').strip()
        doc_type = request.args.get('type', '').strip()
        year_from = request.args.get('year_from', '')
        year_to = request.args.get('year_to', '')

        offset = (page - 1) * per_page

        order_clause = "d.title ASC"
        if sort == 'z_a':
            order_clause = "d.title DESC"
        elif sort == 'old_to_new':
            order_clause = "d.created_at ASC"
        elif sort == 'new_to_old':
            order_clause = "d.created_at DESC"
        elif sort == 'year_asc':
            # SQLite supports NULLS LAST since 3.30.0
            order_clause = "d.publication_year ASC NULLS LAST" if db.engine.name == 'postgresql' or True else "CASE WHEN d.publication_year IS NULL THEN 1 ELSE 0 END, d.publication_year ASC"
        elif sort == 'year_desc':
            order_clause = "d.publication_year DESC NULLS LAST" if db.engine.name == 'postgresql' or True else "CASE WHEN d.publication_year IS NULL THEN 1 ELSE 0 END, d.publication_year DESC"

        conditions = ["1=1"]
        params = {"limit": per_page, "offset": offset}

        if author:
            conditions.append("d.author = :author")
            params['author'] = author

        if search:
            like_op = "ILIKE" if db.engine.name == 'postgresql' else "LIKE"
            conditions.append(f"""
                (d.title {like_op} :search 
                OR d.author {like_op} :search)
            """)
            params['search'] = f"%{search}%"

        if doc_type:
            like_op = "ILIKE" if db.engine.name == 'postgresql' else "LIKE"
            conditions.append(f"d.metadata_type {like_op} :doc_type")
            params['doc_type'] = doc_type

        if year_from:
            conditions.append("d.publication_year >= :year_from")
            params['year_from'] = int(year_from)

        if year_to:
            conditions.append("d.publication_year <= :year_to")
            params['year_to'] = int(year_to)

        where_clause = " AND ".join(conditions)

        count_query = f"SELECT COUNT(*) FROM documents d WHERE {where_clause}"
        total = db.session.execute(text(count_query), params).scalar() or 0

        query = f"""
            SELECT 
                d.id, d.title, d.author, d.metadata_type, d.publication_year, d.date,
                d.page_count, d.volume, d.publisher, d.edition, d.editor,
                d.newspaper_name, d.section, d.column_name, d.publication, d.issue,
                d.created_at, d.updated_at,
                (SELECT file_path FROM document_files WHERE document_id = d.id AND file_type = 'pdf' LIMIT 1) as pdf_path,
                (SELECT CASE WHEN markdown_content IS NOT NULL AND markdown_content != '' THEN true ELSE false END 
                 FROM ocr_results WHERE document_id = d.id LIMIT 1) as has_markdown,
                (SELECT CASE WHEN xml_content IS NOT NULL AND xml_content != '' THEN true ELSE false END 
                 FROM ocr_results WHERE document_id = d.id LIMIT 1) as has_xml,
                (SELECT total_pages FROM ocr_results WHERE document_id = d.id ORDER BY created_at DESC LIMIT 1) as ocr_page_count
            FROM documents d
            WHERE {where_clause}
            ORDER BY {order_clause}
            LIMIT :limit OFFSET :offset
        """

        result = db.session.execute(text(query), params)
        records = []

        for row in result.fetchall():
            display_year = row[4]
            # row[5] is date object or None
            if not display_year and row[5]:
                display_year = extract_year_smart(row[5])
            
            record = {
                "id": str(row[0]),
                "title": row[1] or "Başlıksız",
                "author": row[2],
                "metadata_type": row[3] or "book",
                "publication_year": display_year,
                "date": row[5].isoformat() if row[5] else None,
                "page_count": row[6],
                "volume": row[7],
                "publisher": row[8],
                "edition": row[9],
                "editor": row[10],
                "newspaper_name": row[11],
                "section": row[12],
                "column_name": row[13],
                "journal_name": row[14], 
                "issue_number": row[15], 
                "created_at": row[16].isoformat() if row[16] else None,
                "updated_at": row[17].isoformat() if row[17] else None,
                "pdf_path": row[18],
                "has_markdown": row[19] or False,
                "has_xml": row[20] or False,
                "ocr_page_count": row[21]
            }
            records.append(record)

        return jsonify({
            "records": records,
            "total": total,
            "page": page,
            "per_page": per_page,
            "total_pages": (total + per_page - 1) // per_page
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/records/<record_id>', methods=['GET'])
def get_record(record_id):
    """Tek kayıt detayı"""
    try:
        query = """
            SELECT 
                d.*,
                (SELECT file_path FROM document_files WHERE document_id = d.id AND file_type = 'pdf' LIMIT 1) as pdf_path,
                (SELECT markdown_content FROM ocr_results WHERE document_id = d.id LIMIT 1) as markdown_content,
                (SELECT xml_content FROM ocr_results WHERE document_id = d.id LIMIT 1) as xml_content,
                (SELECT total_pages FROM ocr_results WHERE document_id = d.id LIMIT 1) as ocr_page_count
            FROM documents d
            WHERE d.id = :record_id
        """
        
        result = db.session.execute(text(query), {"record_id": record_id})
        row = result.fetchone()
        
        if not row:
            return jsonify({"error": "Kayıt bulunamadı"}), 404
        
        # Row -> Dict
        try:
            record = dict(row._mapping)
        except AttributeError:
            # Fallback for older SQLAlchemy
            keys = result.keys()
            record = dict(zip(keys, row))
            
        date_fields = ['date', 'access_date', 'created_at', 'updated_at']
        for field in date_fields:
            if record.get(field):
                if hasattr(record[field], 'isoformat'):
                    record[field] = record[field].isoformat()
        
        # UUID fields to string if needed (usually handled by jsonify if UUID obj)
        if hasattr(record['id'], '__str__'):
            record['id'] = str(record['id'])

        record['journal_name'] = record.get('publication')
        record['issue_number'] = record.get('issue')
        
        return jsonify(record)
        
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/records', methods=['POST'])
def create_record():
    """Yeni kayıt oluştur"""
    try:
        data = request.get_json()
        new_id = str(uuid.uuid4())
        now = datetime.now()
        
        record_date = data.get('date') or None
        access_date = data.get('access_date') or None
        pub_year = data.get('publication_year') or None
        
        if record_date and not pub_year:
            try:
                dt = datetime.strptime(record_date, '%Y-%m-%d')
                pub_year = dt.year
            except:
                pass

        query = """
            INSERT INTO documents (
                id, title, author, metadata_type, language, citation_style, url,
                publication_year, date, access_date,
                publisher, publication_city, country, edition, volume, 
                page_count, pages, isbn, series, series_title, series_text, editor,
                publication, issue, doi, issn, journal_abbreviation,
                newspaper_name, publication_place, section, column_name, page_range,
                encyclopedia_title, short_title,
                archive, archive_location, library_catalog, call_number, rights,
                created_at, updated_at
            ) VALUES (
                :id, :title, :author, :metadata_type, :language, :citation_style, :url,
                :publication_year, :date, :access_date,
                :publisher, :publication_city, :country, :edition, :volume, 
                :page_count, :pages, :isbn, :series, :series_title, :series_text, :editor,
                :publication, :issue, :doi, :issn, :journal_abbreviation,
                :newspaper_name, :publication_place, :section, :column_name, :page_range,
                :encyclopedia_title, :short_title,
                :archive, :archive_location, :library_catalog, :call_number, :rights,
                :created_at, :updated_at
            )
        """
        
        params = {
            "id": new_id,
            "title": data.get('title', 'Başlıksız'),
            "author": data.get('author'),
            "metadata_type": data.get('metadata_type', 'book'),
            "language": data.get('language'),
            "citation_style": data.get('citation_style'),
            "url": data.get('url'),
            "publication_year": pub_year,
            "date": record_date,
            "access_date": access_date,
            "publisher": data.get('publisher'),
            "publication_city": data.get('publication_city'),
            "country": data.get('country'),
            "edition": data.get('edition'),
            "volume": data.get('volume'),
            "page_count": data.get('page_count'),
            "pages": data.get('page_count'),
            "isbn": data.get('isbn'),
            "series": data.get('series'),
            "series_title": data.get('series_title'),
            "series_text": data.get('series_text'),
            "editor": data.get('editor'),
            "publication": data.get('publication'),
            "issue": data.get('issue'),
            "doi": data.get('doi'),
            "issn": data.get('issn'),
            "journal_abbreviation": data.get('journal_abbreviation'),
            "newspaper_name": data.get('newspaper_name'),
            "publication_place": data.get('publication_place'),
            "section": data.get('section'),
            "column_name": data.get('column_name'),
            "page_range": data.get('page_range'),
            "encyclopedia_title": data.get('encyclopedia_title'),
            "short_title": data.get('short_title'),
            "archive": data.get('archive'),
            "archive_location": data.get('archive_location'),
            "library_catalog": data.get('library_catalog'),
            "call_number": data.get('call_number'),
            "rights": data.get('rights'),
            "created_at": now,
            "updated_at": now
        }
        
        db.session.execute(text(query), params)
        db.session.commit()
        
        return jsonify({"success": True, "id": new_id, "message": "Kayıt başarıyla oluşturuldu"}), 201
        
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/records/<record_id>', methods=['PUT'])
def update_record(record_id):
    """Kayıt güncelle"""
    try:
        data = request.get_json()
        
        updatable_fields = [
            'title', 'author', 'metadata_type', 'language', 'country', 'citation_style', 'url',
            'publication_year', 'date', 'access_date',
            'publisher', 'publication_city', 'edition', 'volume', 
            'page_count', 'pages', 'isbn', 'series', 'series_title', 'series_text', 'editor',
            'publication', 'issue', 'doi', 'issn', 'journal_abbreviation',
            'newspaper_name', 'publication_place', 'section', 'column_name', 'page_range',
            'encyclopedia_title', 'short_title',
            'archive', 'archive_location', 'library_catalog', 'call_number', 'rights'
        ]
        
        set_clauses = ["updated_at = :now"]
        params = {"record_id": record_id, "now": datetime.now()}
        
        for field in updatable_fields:
            val = None
            if field == 'publication_city' and 'city' in data: val = data.get('city')
            elif field == 'pages' and 'page_count' in data: val = data.get('page_count')
            elif field in data:
                val = data[field]
            
            if val == '': val = None
                
            if field in data or val is not None:
                set_clauses.append(f"{field} = :{field}")
                params[field] = val
        
        query = f"""
            UPDATE documents
            SET {', '.join(set_clauses)}
            WHERE id = :record_id
        """
        
        db.session.execute(text(query), params)
        db.session.commit()
        
        return jsonify({"success": True, "message": "Kayıt başarıyla güncellendi"})
        
    except Exception as e:
        db.session.rollback()
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e), "traceback": traceback.format_exc()}), 500

@admin_bp.route('/api/records/<record_id>', methods=['DELETE'])
def delete_record(record_id):
    """Kayıt sil (hard delete)"""
    try:
        file_result = db.session.execute(text("""
            SELECT file_path FROM document_files WHERE document_id = :record_id
        """), {"record_id": record_id})
        file_paths = [row[0] for row in file_result.fetchall()]
        
        db.session.execute(text("DELETE FROM ocr_results WHERE document_id = :record_id"), 
                   {"record_id": record_id})
        db.session.execute(text("DELETE FROM document_files WHERE document_id = :record_id"), 
                   {"record_id": record_id})
        
        result = db.session.execute(text("DELETE FROM documents WHERE id = :record_id"), 
                            {"record_id": record_id})
        
        db.session.commit()
        
        if result.rowcount == 0:
            return jsonify({"error": "Kayıt bulunamadı"}), 404
        
        
        # Dosya silme işlemleri (PDF, MD, XML)
        deleted_count = 0
        
        # 1. document_files tablosundaki asıl dosyaları sil
        for file_path in file_paths:
            try:
                if file_path:
                    path_obj = Path(file_path)
                    if path_obj.exists():
                        path_obj.unlink()
                        deleted_count += 1
            except Exception:
                pass 
        
        # 2. ocr_results tablosundan dolaylı oluşturulan dosyaları temizlemeyi deneyelim
        # (MD ve XML içerikleri veritabanında saklansa da, diskte de 'data/outputs' altında olabilirler)
        # Ancak current implementation sadece DB'de saklıyor gibi görünüyor.
        # Yine de oluşturulmuş olabilecek .md ve .xml dosyalarını dosya adından tahmin edip silebiliriz.
        # Şimdilik sadece kayıtlı file_paths siliniyor, bu yeterli.
        
        return jsonify({
            "success": True,
            "message": "Kayıt ve ilişkili dosyalar silindi",
            "deleted_files": deleted_count
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500

# ==========================================
# SEARCH, CONTENT handlers
# ==========================================

@admin_bp.route('/api/search', methods=['GET'])
def search_records():
    """Gelişmiş arama"""
    try:
        query_text = request.args.get('q', '').strip()
        search_field = request.args.get('field', 'all')
        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 25))
        
        if not query_text:
            return jsonify({"records": [], "total": 0})
        
        offset = (page - 1) * per_page
        conditions = []
        params = {"limit": per_page, "offset": offset}
        
        like_op = "ILIKE" if db.engine.name == 'postgresql' else "LIKE"
        if search_field == 'title':
            conditions.append(f"d.title {like_op} :search")
            params['search'] = f"%{query_text}%"
        elif search_field == 'author':
            conditions.append(f"d.author {like_op} :search")
            params['search'] = f"%{query_text}%"
        elif search_field == 'year':
            try:
                year = int(query_text)
                conditions.append("d.publication_year = :year")
                params['year'] = year
            except ValueError:
                return jsonify({"error": "Geçersiz yıl"}), 400
        else:
            conditions.append(f"""
                (d.title {like_op} :search 
                OR d.author {like_op} :search
                OR CAST(d.publication_year AS TEXT) LIKE :search)
            """)
            params['search'] = f"%{query_text}%"
        
        where_clause = " AND ".join(conditions) if conditions else "1=1"
        
        count_query = f"SELECT COUNT(*) FROM documents d WHERE {where_clause}"
        total = db.session.execute(text(count_query), params).scalar() or 0
        
        query = f"""
            SELECT 
                d.id, d.title, d.author, d.metadata_type,
                d.publication_year, d.created_at,
                (SELECT file_path FROM document_files WHERE document_id = d.id AND file_type = 'pdf' LIMIT 1) as pdf_path
            FROM documents d
            WHERE {where_clause}
            ORDER BY d.title ASC
            LIMIT :limit OFFSET :offset
        """
        
        result = db.session.execute(text(query), params)
        records = [
            {
                "id": str(row[0]),
                "title": row[1],
                "author": row[2],
                "metadata_type": row[3],
                "publication_year": row[4],
                "created_at": row[5].isoformat() if row[5] else None,
                "pdf_path": row[6]
            }
            for row in result.fetchall()
        ]
        
        return jsonify({
            "records": records,
            "total": total,
            "page": page,
            "per_page": per_page
        })
        
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ==========================================
# FILE & CONTENT HANDLERS
# ==========================================

@admin_bp.route('/api/records/<record_id>/markdown', methods=['GET'])
def get_markdown(record_id):
    try:
        result = db.session.execute(text("SELECT markdown_content FROM ocr_results WHERE document_id = :id"), {"id": record_id})
        row = result.fetchone()
        return jsonify({"content": row[0] if row else ""})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/records/<record_id>/markdown', methods=['PUT'])
def save_markdown(record_id):
    try:
        data = request.get_json()
        content = data.get('content', '')
        
        check_query = "SELECT id FROM ocr_results WHERE document_id = :doc_id LIMIT 1"
        existing = db.session.execute(text(check_query), {"doc_id": record_id}).fetchone()
        
        if existing:
            update_query = "UPDATE ocr_results SET markdown_content = :content WHERE document_id = :doc_id"
            db.session.execute(text(update_query), {"doc_id": record_id, "content": content})
        else:
            insert_query = """
                INSERT INTO ocr_results (id, document_id, markdown_content, created_at, total_pages)
                VALUES (:id, :doc_id, :content, :now, 0)
            """
            db.session.execute(text(insert_query), {
                "id": str(uuid.uuid4()),
                "doc_id": record_id,
                "content": content,
                "now": datetime.now()
            })
            
        db.session.commit()
        return jsonify({"success": True, "message": "Markdown kaydedildi"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/records/<record_id>/xml', methods=['GET'])
def get_xml(record_id):
    try:
        result = db.session.execute(text("SELECT xml_content FROM ocr_results WHERE document_id = :id"), {"id": record_id})
        row = result.fetchone()
        return jsonify({"content": row[0] if row else ""})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/records/<record_id>/xml', methods=['PUT'])
def save_xml(record_id):
    try:
        data = request.get_json()
        content = data.get('content', '')
        
        check_query = "SELECT id FROM ocr_results WHERE document_id = :doc_id LIMIT 1"
        existing = db.session.execute(text(check_query), {"doc_id": record_id}).fetchone()
        
        if existing:
            update_query = "UPDATE ocr_results SET xml_content = :content WHERE document_id = :doc_id"
            db.session.execute(text(update_query), {"doc_id": record_id, "content": content})
        else:
            insert_query = """
                INSERT INTO ocr_results (id, document_id, xml_content, created_at, total_pages)
                VALUES (:id, :doc_id, :content, :now, 0)
            """
            db.session.execute(text(insert_query), {
                "id": str(uuid.uuid4()),
                "doc_id": record_id,
                "content": content,
                "now": datetime.now()
            })
            
        db.session.commit()
        return jsonify({"success": True, "message": "XML kaydedildi"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/files/download/<record_id>')
def download_source_file(record_id):
    """Kaynak dosyayı (PDF/JPG) indir - Düzenli isimlendirme ile"""
    try:
        file_query = text("SELECT file_path FROM document_files WHERE document_id = :record_id LIMIT 1")
        file_res = db.session.execute(file_query, {"record_id": record_id}).fetchone()
        
        if not file_res or not file_res[0]:
            return jsonify({"error": "Dosya bulunamadı"}), 404
            
        base_dirs = [
            Path('/app/data/veriler'),
            Path(current_app.root_path).parent / 'app' / 'database' / 'veriler',
            Path(current_app.root_path) / 'database' / 'veriler'
        ]
        if current_app.config.get('UPLOADS_DIR'):
            base_dirs.append(Path(current_app.config.get('UPLOADS_DIR')))
            
        file_path = resolve_file_path(file_res[0], base_dirs)
        if not file_path:
            return jsonify({"error": "Fiziksel dosya diskte yok"}), 404

        doc_query = text("SELECT * FROM documents WHERE id = :id")
        doc_row = db.session.execute(doc_query, {"id": record_id}).fetchone()
        
        download_name = None
        if doc_row:
            try:
                record = dict(doc_row._mapping)
            except AttributeError:
                record = dict(doc_row)
            
            ext = os.path.splitext(file_path)[1]
            safe_title = generate_detailed_filename(record)
            download_name = f"{safe_title}{ext}"

        return send_file(
            file_path, 
            as_attachment=True,
            download_name=download_name
        )
        
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/content/download/<record_id>/<content_type>')
def download_generated_content(record_id, content_type):
    """Markdown veya XML içeriğini dosya olarak indir"""
    try:
        query = text("SELECT * FROM documents WHERE id = :id")
        row = db.session.execute(query, {"id": record_id}).fetchone()
        
        if not row:
            return jsonify({"error": "Kayıt bulunamadı"}), 404

        try:
            record = dict(row._mapping)
        except AttributeError:
            record = dict(row)

        safe_title = generate_detailed_filename(record)

        col_name = "markdown_content" if content_type == "markdown" else "xml_content"
        ext = "md" if content_type == "markdown" else "xml"
        
        query = text(f"SELECT {col_name} FROM ocr_results WHERE document_id = :id")
        result = db.session.execute(query, {"id": record_id}).fetchone()
        
        content = result[0] if result and result[0] is not None else ""
        
        if not content and not result:
            # Sadece hiç sonuç yoksa 404 dönüyoruz, içerik string olarak boşsa 404 dönmeye gerek yok.
            return jsonify({"error": "İçerik bulunamadı"}), 404

        content_bytes = content.encode('utf-8')
        
        return Response(
            content_bytes,
            mimetype='text/plain',
            headers={
                'Content-Disposition': f'attachment; filename="{safe_title}.{ext}"',
                'Content-Length': str(len(content_bytes))
            }
        )
        
    except Exception as e:
        print(f"Content Download Error: {e}")
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/files/view/<record_id>')
def view_file_admin(record_id):
    try:
        result = db.session.execute(text("SELECT file_path, file_type FROM document_files WHERE document_id = :id AND file_type = 'pdf' LIMIT 1"), {"id": record_id})
        row = result.fetchone()
        if not row:
            return "File not found", 404
            
        file_path = row[0]
        base_dirs = [
            Path('/app/data/veriler'),
            Path(current_app.root_path).parent / 'app' / 'database' / 'veriler',
            Path(current_app.root_path) / 'database' / 'veriler'
        ]
        if current_app.config.get('UPLOADS_DIR'):
            base_dirs.append(Path(current_app.config.get('UPLOADS_DIR')))
            
        resolved = resolve_file_path(file_path, base_dirs)
        
        target_path = resolved if resolved else file_path
        
        if not os.path.exists(target_path):
            return f"File path does not exist on server: {target_path}", 404
            
        return send_from_directory(os.path.dirname(target_path), os.path.basename(target_path))
    except Exception as e:
        return str(e), 500

@admin_bp.route('/api/export/record/<record_id>/zip', methods=['GET'])
def export_record_zip(record_id):
    """Bir esere ait her şeyi (PDF, MD, XML, JSON) ZIP olarak indir"""
    try:
        query = text("""
            SELECT 
                d.*,
                (SELECT file_path FROM document_files WHERE document_id = d.id LIMIT 1) as file_path,
                (SELECT markdown_content FROM ocr_results WHERE document_id = d.id LIMIT 1) as markdown,
                (SELECT xml_content FROM ocr_results WHERE document_id = d.id LIMIT 1) as xml
            FROM documents d
            WHERE d.id = :id
        """)
        
        row = db.session.execute(query, {"id": record_id}).fetchone()
        
        if not row:
            return jsonify({"error": "Kayıt bulunamadı"}), 404

        try:
            record = dict(row._mapping)
        except AttributeError:
            record = dict(row)
        
        safe_title = generate_detailed_filename(record)
        
        memory_file = BytesIO()
        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            
            # A. Metadata JSON
            metadata = {k: v for k, v in record.items() if k not in ['markdown', 'xml', 'file_path']}
            for k, v in metadata.items():
                if isinstance(v, uuid.UUID):
                    metadata[k] = str(v)
                elif hasattr(v, 'isoformat'): 
                    metadata[k] = v.isoformat()
            
            zf.writestr(f"{safe_title}_metadata.json", json.dumps(metadata, ensure_ascii=False, indent=2))
            
            # B. Markdown
            if record.get('markdown'):
                zf.writestr(f"{safe_title}.md", record.get('markdown'))
                
            # C. XML
            if record.get('xml'):
                zf.writestr(f"{safe_title}.xml", record.get('xml'))
                
            # D. Kaynak Dosya
            if record.get('file_path'):
                base_dirs = [
                    Path('/app/data/veriler'),
                    Path(current_app.root_path).parent / 'app' / 'database' / 'veriler',
                    Path(current_app.root_path) / 'database' / 'veriler'
                ]
                if current_app.config.get('UPLOADS_DIR'):
                    base_dirs.append(Path(current_app.config.get('UPLOADS_DIR')))
                    
                real_path = resolve_file_path(record['file_path'], base_dirs)
                if real_path and os.path.exists(real_path):
                    ext = os.path.splitext(real_path)[1]
                    zf.write(real_path, arcname=f"{safe_title}{ext}")

        memory_file.seek(0)
        zip_content = memory_file.getvalue()
        
        return Response(
            zip_content,
            mimetype='application/zip',
            headers={
                'Content-Disposition': f'attachment; filename="{safe_title}.zip"',
                'Content-Length': str(len(zip_content))
            }
        )

    except Exception as e:
        print(f"ZIP Export Error: {e}")
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/export/author/<author_name>/zip', methods=['GET'])
def export_author_zip(author_name):
    """Bir yazara ait tüm eserleri ZIP olarak indir"""
    try:
        # 1. Yazara ait tüm kayıtları bul
        query = text("""
            SELECT 
                d.*,
                (SELECT file_path FROM document_files WHERE document_id = d.id LIMIT 1) as file_path,
                (SELECT markdown_content FROM ocr_results WHERE document_id = d.id LIMIT 1) as markdown,
                (SELECT xml_content FROM ocr_results WHERE document_id = d.id LIMIT 1) as xml
            FROM documents d
            WHERE d.author = :author
        """)
        
        rows = db.session.execute(query, {"author": author_name}).fetchall()
        
        if not rows:
            return jsonify({"error": "Yazara ait kayıt bulunamadı"}), 404

        memory_file = BytesIO()
        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            
            # Her bir kayıt için dosyaları ekle
            for row in rows:
                try:
                    record = dict(row._mapping)
                except AttributeError:
                    record = dict(row)
                
                # Eser başlığına göre klasör adı oluştur
                safe_title = generate_detailed_filename(record)
                folder_prefix = f"{safe_title}/"
                
                # A. Metadata JSON
                metadata = {k: v for k, v in record.items() if k not in ['markdown', 'xml', 'file_path']}
                for k, v in metadata.items():
                    if isinstance(v, uuid.UUID):
                        metadata[k] = str(v)
                    elif hasattr(v, 'isoformat'): 
                        metadata[k] = v.isoformat()
                
                zf.writestr(f"{folder_prefix}{safe_title}_metadata.json", json.dumps(metadata, ensure_ascii=False, indent=2))
                
                # B. Markdown
                if record.get('markdown'):
                    zf.writestr(f"{folder_prefix}{safe_title}.md", record.get('markdown'))
                    
                # C. XML
                if record.get('xml'):
                    zf.writestr(f"{folder_prefix}{safe_title}.xml", record.get('xml'))
                    
                # D. Kaynak Dosya
                if record.get('file_path'):
                    base_dirs = [
                        Path('/app/data/veriler'),
                        Path(current_app.root_path).parent / 'app' / 'database' / 'veriler',
                        Path(current_app.root_path) / 'database' / 'veriler'
                    ]
                    if current_app.config.get('UPLOADS_DIR'):
                        base_dirs.append(Path(current_app.config.get('UPLOADS_DIR')))
                        
                    real_path = resolve_file_path(record['file_path'], base_dirs)
                    if real_path and os.path.exists(real_path):
                        ext = os.path.splitext(real_path)[1]
                        zf.write(real_path, arcname=f"{folder_prefix}{safe_title}{ext}")

        memory_file.seek(0)
        zip_content = memory_file.getvalue()
        
        filename = f"{secure_filename_tr(author_name)}_Arsiv.zip"
        
        return Response(
            zip_content,
            mimetype='application/zip',
            headers={
                'Content-Disposition': f'attachment; filename="{filename}"',
                'Content-Length': str(len(zip_content))
            }
        )

    except Exception as e:
        print(f"Author ZIP Export Error: {e}")
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/export/json', methods=['GET'])
def export_json():
    """Tüm veritabanını JSON olarak dışa aktar"""
    try:
        from app.models import Document
        
        author_filter = request.args.get('author', '').strip()
        
        query = db.session.query(Document)
        if author_filter:
            query = query.filter(Document.author == author_filter)
            
        documents = query.order_by(Document.author, Document.title).all()
        
        records_list = []
        for doc in documents:
            # Get OCR contents if they exist
            ocr = doc.ocr_results[0] if doc.ocr_results else None
            
            record_dict = {
                "id": str(doc.id),
                "title": doc.title,
                "author": doc.author,
                "metadata_type": doc.metadata_type,
                "publication_year": doc.publication_year,
                "page_count": doc.page_count,
                "volume": doc.volume,
                "publisher": doc.publisher,
                "edition": doc.edition,
                "editor": doc.editor,
                "description": getattr(doc, 'description', None), # Check if exists safely
                "language": doc.language,
                "country": doc.country,
                "publication_city": doc.publication_city,
                "isbn": doc.isbn,
                "publication": doc.publication,
                "issue": doc.issue,
                "doi": doc.doi,
                "issn": doc.issn,
                "newspaper_name": doc.newspaper_name,
                "access_date": doc.access_date.isoformat() if doc.access_date else None,
                "date": doc.date.isoformat() if doc.date else None,
                "created_at": doc.created_at.isoformat() if doc.created_at else None,
                "markdown_content": ocr.markdown_content if ocr else None,
                "xml_content": ocr.xml_content if ocr else None,
                "ocr_total_pages": ocr.total_pages if ocr else 0
            }
            records_list.append(record_dict)
            
        export_data = {
            "export_date": datetime.now().isoformat(),
            "total_records": len(records_list),
            "records": records_list
        }
        
        filename = f"database_export_{datetime.now().strftime('%Y%p%d_%H%M%S')}.json"
        
        return Response(
            json.dumps(export_data, ensure_ascii=False, indent=2),
            mimetype='application/json',
            headers={'Content-Disposition': f'attachment; filename={filename}'}
        )
    except Exception as e:
        print(f"JSON Export Error: {e}")
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/export/csv', methods=['GET'])
def export_csv():
    """Veritabanını CSV olarak dışa aktar"""
    try:
        import csv
        from io import StringIO
        from app.models import Document
        
        documents = db.session.query(Document).order_by(Document.author, Document.title).all()
        
        output = StringIO()
        writer = csv.writer(output)
        
        headers = [
            'ID', 'Başlık', 'Yazar', 'Tür', 'Tarih / Yıl',
            'Sayfa Sayısı', 'Cilt', 'Yayınevi', 'Baskı',
            'Editör', 'Dil', 'Oluşturma Tarihi'
        ]
        writer.writerow(headers)
        
        for doc in documents:
            display_date = ""
            if doc.date:
                display_date = doc.date.isoformat()
            elif doc.publication_year:
                display_date = str(doc.publication_year)
                
            writer.writerow([
                str(doc.id),
                doc.title,
                doc.author or '',
                doc.metadata_type,
                display_date,
                doc.page_count or '',
                doc.volume or '',
                doc.publisher or '',
                doc.edition or '',
                doc.editor or '',
                doc.language or '',
                doc.created_at.isoformat() if doc.created_at else ''
            ])
            
        filename = f"database_export_{datetime.now().strftime('%Y%p%d_%H%M%S')}.csv"
        
        return Response(
            output.getvalue(),
            mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename={filename}'}
        )
    except Exception as e:
        print(f"CSV Export Error: {e}")
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/export/excel', methods=['GET'])
def export_excel():
    """Veritabanını Excel (.xlsx) olarak dışa aktar"""
    try:
        import openpyxl
        from openpyxl.styles import Font
        from app.models import Document
        
        documents = db.session.query(Document).order_by(Document.author, Document.title).all()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kayıtlar"
        
        headers = [
            'ID', 'Başlık', 'Yazar', 'Tür', 'Tarih / Yıl',
            'Sayfa Sayısı', 'Cilt', 'Yayınevi', 'Baskı',
            'Editör', 'Dil', 'Oluşturma Tarihi'
        ]
        ws.append(headers)
        
        for cell in ws[1]:
            cell.font = Font(bold=True)
            
        for doc in documents:
            display_date = ""
            if doc.date:
                display_date = doc.date.isoformat()
            elif doc.publication_year:
                display_date = str(doc.publication_year)
                
            ws.append([
                str(doc.id),
                doc.title,
                doc.author or '',
                doc.metadata_type,
                display_date,
                doc.page_count or '',
                doc.volume or '',
                doc.publisher or '',
                doc.edition or '',
                doc.editor or '',
                doc.language or '',
                doc.created_at.isoformat() if doc.created_at else ''
            ])
            
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"database_export_{datetime.now().strftime('%Y%p%d_%H%M%S')}.xlsx"
        
        return Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={filename}'}
        )
    except Exception as e:
        print(f"Excel Export Error: {e}")
        return jsonify({"error": str(e)}), 500

@admin_bp.route('/api/import', methods=['POST'])
def import_data():
    """JSON dosyasından veri içe aktar"""
    try:
        from app.models import Document, OCRResult
        
        if 'file' not in request.files:
            return jsonify({"error": "Dosya seçilmedi"}), 400
            
        file = request.files['file']
        if file.filename == '':
            return jsonify({"error": "Dosya seçilmedi"}), 400
            
        if not file.filename.endswith('.json'):
            return jsonify({"error": "Sadece JSON dosyaları desteklenir"}), 400
            
        data = json.load(file)
        records = data.get('records', [])
        
        imported = 0
        skipped = 0
        errors = []
        
        for rec in records:
            try:
                # Check for existing record
                rec_id = rec.get('id')
                if rec_id:
                    exists = db.session.query(Document).filter(Document.id == rec_id).first()
                    if exists:
                        skipped += 1
                        continue
                
                # Create new document
                doc = Document(
                    title=rec.get('title', 'Başlıksız'),
                    author=rec.get('author'),
                    metadata_type=rec.get('metadata_type', 'book'),
                    publication_year=rec.get('publication_year'),
                    page_count=rec.get('page_count'),
                    volume=rec.get('volume'),
                    publisher=rec.get('publisher'),
                    edition=rec.get('edition'),
                    editor=rec.get('editor'),
                    language=rec.get('language', 'tr'),
                    country=rec.get('country'),
                    publication_city=rec.get('publication_city'),
                    isbn=rec.get('isbn'),
                    publication=rec.get('publication'),
                    issue=rec.get('issue'),
                    doi=rec.get('doi'),
                    issn=rec.get('issn'),
                    newspaper_name=rec.get('newspaper_name')
                )
                
                if rec_id:
                    doc.id = rec_id
                
                if rec.get('date'):
                    try: doc.date = datetime.fromisoformat(rec['date']).date()
                    except: pass
                
                if rec.get('access_date'):
                    try: doc.access_date = datetime.fromisoformat(rec['access_date']).date()
                    except: pass
                
                db.session.add(doc)
                db.session.flush() # Get ID if not provided
                
                # Add OCR results if present
                if rec.get('markdown_content') or rec.get('xml_content'):
                    # 1. Öncelikli olarak XML içerisindeki <sayfa ...> veya <page ...> etiketlerini say
                    ocr_pages = 0
                    if rec.get('xml_content'):
                        import re
                        xml_text = rec.get('xml_content', '')
                        matches = re.findall(r'<(?:sayfa|page)\b', xml_text, re.IGNORECASE)
                        if matches:
                            ocr_pages = len(matches)
                            
                    # 2. XML'den bulunamadıysa (veya XML yoksa), JSON export'taki 'ocr_total_pages' değerine bak
                    if not ocr_pages:
                        ocr_pages = rec.get('ocr_total_pages')
                    
                    # 3. Hala yoksa document metadata'sındaki 'page_count'u kullan
                    if not ocr_pages:
                        try:
                            ocr_pages = int(rec.get('page_count') or 0)
                        except (ValueError, TypeError):
                            ocr_pages = 0
                            
                    # Her şeye rağmen hala 0 veya boş ise, en az 1 yapalım
                    if not ocr_pages:
                        ocr_pages = 1
                        
                    ocr = OCRResult(
                        document_id=doc.id,
                        markdown_content=rec.get('markdown_content'),
                        xml_content=rec.get('xml_content'),
                        total_pages=ocr_pages
                    )
                    db.session.add(ocr)
                
                imported += 1
            except Exception as e:
                errors.append(f"Hata ({rec.get('title')}): {str(e)}")
                
        db.session.commit()
        return jsonify({
            "success": True,
            "imported": imported,
            "skipped": skipped,
            "errors": errors[:10]
        })
    except Exception as e:
        db.session.rollback()
        print(f"Import Error: {e}")
        return jsonify({"error": str(e)}), 500
